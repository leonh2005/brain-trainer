"""產出前 300 家清單（含賣場首頁查到的統編與公司名），不篩除，供人工核對。
輸出 Excel + HTML。"""
import html
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = "/tmp/shop_taxid_top300.json"
XLSX = Path.home() / "Desktop" / "蝦皮前300大_含統編.xlsx"
HTM = Path("/Users/steven/CCProject/shopee_top300_taxid.html")

shops = json.load(open(SRC))
# 依「是否查到統編」再依月營業額排序：有名單的集中在前面方便核對
shops.sort(key=lambda s: (0 if s.get("shop_taxid") else 1, -s["month_rev"]))
n_have = sum(1 for s in shops if s.get("shop_taxid"))

# ---------- Excel ----------
HEAD = PatternFill("solid", fgColor="1F3864")
HF = Font(color="FFFFFF", bold=True)
wb = Workbook()
ws = wb.active
ws.title = "前300大_含統編"
heads = ["#", "賣場名稱", "月營業額估", "單品最大", "件數", "地區", "蝦皮帳號",
         "賣場統編", "公司名稱", "最高單品"]
ws.append(heads)
for c in range(1, len(heads) + 1):
    ws.cell(row=1, column=c).fill, ws.cell(row=1, column=c).font = HEAD, HF
    ws.cell(row=1, column=c).alignment = Alignment(horizontal="center")
for i, s in enumerate(shops, 1):
    ws.append([i, s["name"], int(s["month_rev"]), int(s["max_item_rev"]), s["n"],
               s.get("loc", ""), s.get("account") or "", s.get("shop_taxid") or "",
               s.get("shop_company") or "", s.get("max_item") or ""])
for w, col in zip([5, 42, 14, 14, 8, 13, 20, 12, 26, 38], range(1, 11)):
    ws.column_dimensions[get_column_letter(col)].width = w
for r in range(2, ws.max_row + 1):
    for c in (3, 4):
        ws.cell(row=r, column=c).number_format = "#,##0"
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:J{ws.max_row}"
wb.save(XLSX)

# ---------- HTML ----------
esc = html.escape


def row(s, i):
    have = bool(s.get("shop_taxid"))
    cls = "have" if have else ""
    return (f"<tr class='{cls}'>"
            f"<td class='ck'><input type='checkbox' data-k=\"{esc(s['name'])}\" onchange='mark(this)'></td>"
            f"<td class='r'>{i}</td>"
            f"<td class='name'>{esc(s['name'])}</td>"
            f"<td class='r'>{int(s['month_rev']):,}</td>"
            f"<td class='r'>{int(s['max_item_rev']):,}</td>"
            f"<td class='r'>{s['n']}</td>"
            f"<td>{esc(s.get('loc') or '')}</td>"
            f"<td class='mono'>{esc(s.get('account') or '')}</td>"
            f"<td class='mono tax'>{esc(s.get('shop_taxid') or '')}</td>"
            f"<td>{esc(s.get('shop_company') or '')}</td>"
            f"</tr>")


rows = "\n".join(row(s, i + 1) for i, s in enumerate(shops))
doc = f"""<!doctype html>
<html lang="zh-Hant"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>蝦皮前300大 含統編</title>
<style>
:root{{--bg:#fff;--fg:#1a1a1a;--muted:#6b7280;--line:#e5e7eb;--accent:#0f766e;--ok:#047857}}
@media(prefers-color-scheme:dark){{:root{{--bg:#111214;--fg:#e8e8ea;--muted:#9ca3af;--line:#2a2c30;--accent:#5eead4;--ok:#34d399}}}}
*{{box-sizing:border-box}}
body{{margin:0;background:var(--bg);color:var(--fg);font:15px/1.6 -apple-system,"PingFang TC","Noto Sans TC",sans-serif}}
.wrap{{max-width:1280px;margin:0 auto;padding:30px 18px 60px}}
h1{{font-size:23px;margin:0 0 4px}} .sub{{color:var(--muted);font-size:13px;margin-bottom:18px}}
.stats{{display:flex;gap:12px;flex-wrap:wrap;margin-bottom:18px}}
.card{{border:1px solid var(--line);border-radius:10px;padding:12px 16px}}
.card .n{{font-size:22px;font-weight:700;color:var(--accent)}} .card .l{{font-size:12px;color:var(--muted)}}
input#q{{width:100%;padding:10px 12px;border:1px solid var(--line);border-radius:8px;background:var(--bg);color:var(--fg);margin-bottom:14px}}
table{{width:100%;border-collapse:collapse;font-size:13px}}
th,td{{padding:7px 9px;border-bottom:1px solid var(--line);text-align:left}}
th{{position:sticky;top:0;background:var(--bg);cursor:pointer;font-size:12px;color:var(--muted);white-space:nowrap}}
td.r,th.r{{text-align:right;font-variant-numeric:tabular-nums}}
td.name{{max-width:280px;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}}
td.mono{{font-variant-numeric:tabular-nums;font-family:ui-monospace,monospace}}
td.tax{{color:var(--ok);font-weight:600}}
tr.have{{background:color-mix(in srgb,var(--ok) 7%,transparent)}}
.ck{{width:30px}} .ck input{{width:auto;cursor:pointer}}
tr.done{{opacity:.35}} tr.done td.name{{text-decoration:line-through}}
</style></head><body><div class="wrap">
<h1>蝦皮前 300 大賣場 — 含賣場首頁統編</h1>
<div class="sub">2026-10-07 · 未揭露統編且非官方/優選，按月營業額估排序前 300 · 統編取自各賣場首頁（company_id）· 未做篩除，供人工核對</div>
<div class="stats">
<div class="card"><div class="n">{len(shops)}</div><div class="l">總數</div></div>
<div class="card"><div class="n">{n_have}</div><div class="l">賣場首頁查到統編（綠底）</div></div>
<div class="card"><div class="n">{len(shops)-n_have}</div><div class="l">查無統編</div></div>
</div>
<input id="q" placeholder="搜尋賣場名、帳號、公司名、統編…">
<table id="t"><thead><tr>
<th class="ck">✓</th><th class="r" data-k="1">#</th><th data-k="2">賣場名稱</th>
<th class="r" data-k="3">月營業額估</th><th class="r" data-k="4">單品最大</th>
<th class="r" data-k="5">件數</th><th data-k="6">地區</th><th data-k="7">蝦皮帳號</th>
<th data-k="8">賣場統編</th><th data-k="9">公司名稱</th>
</tr></thead><tbody>
{rows}
</tbody></table>
</div>
<script>
const CK='shopee_top300_checked_v1';
let ck=new Set(JSON.parse(localStorage.getItem(CK)||'[]'));
function applyAll(){{document.querySelectorAll('input[data-k]').forEach(cb=>{{
  const on=ck.has(cb.dataset.k);cb.checked=on;cb.closest('tr').classList.toggle('done',on);}});}}
function mark(cb){{const k=cb.dataset.k;
  if(cb.checked)ck.add(k);else ck.delete(k);
  cb.closest('tr').classList.toggle('done',cb.checked);
  localStorage.setItem(CK,JSON.stringify([...ck]));}}
document.getElementById('q').addEventListener('input',e=>{{
  const v=e.target.value.trim().toLowerCase();
  document.querySelectorAll('#t tbody tr').forEach(r=>{{
    r.style.display=r.textContent.toLowerCase().includes(v)?'':'none';}});}});
{{ const t=document.getElementById('t'); let dir={{}};
  t.tHead.addEventListener('click',e=>{{const th=e.target.closest('th');if(!th||!th.dataset.k)return;
    const k=+th.dataset.k;dir[k]=!dir[k];const s=dir[k]?1:-1;
    const num=[1,3,4,5].includes(k);
    const rows=[...t.tBodies[0].rows].sort((a,b)=>{{let x=a.cells[k].textContent,y=b.cells[k].textContent;
      if(num){{x=+x.replace(/,/g,'')||0;y=+y.replace(/,/g,'')||0;return (x-y)*s;}}
      return x.localeCompare(y,'zh-Hant')*s;}});
    rows.forEach(r=>t.tBodies[0].appendChild(r));}}); }}
applyAll();
</script></body></html>"""
HTM.write_text(doc, encoding="utf-8")
print("Excel:", XLSX)
print("HTML :", HTM)
print(f"共 {len(shops)} 家，{n_have} 家查到統編")
